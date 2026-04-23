import os
import asyncio
from openpyxl import Workbook, load_workbook
!pip install aiogram openpyxl nest_asyncio

from aiogram import Bot, Dispatcher
from aiogram.types import Message, ReplyKeyboardMarkup, KeyboardButton, ReplyKeyboardRemove
from aiogram.filters import CommandStart
from aiogram.fsm.state import StatesGroup, State
from aiogram.fsm.context import FSMContext
from aiogram.fsm.storage.memory import MemoryStorage

# --- SOZLAMALAR ---
TOKEN = "8794870759:AAE9FlpYQSDGHIzvKgUozdP39BzEXawEieU"
ADMIN_ID = 5662077584

bot = Bot(token=TOKEN)
dp = Dispatcher(storage=MemoryStorage())

# admin reply uchun mapping
msg_map = {}

# --- EXCEL ---
def save_to_excel(user_id, username, phone, lang, category, text, file_id):
    file_name = "data.xlsx"

    if not os.path.exists(file_name):
        wb = Workbook()
        ws = wb.active
        ws.append(["UserID", "Username", "Phone", "Lang", "Category", "Text", "FileID"])
        wb.save(file_name)

    wb = load_workbook(file_name)
    ws = wb.active
    ws.append([user_id, username, phone, lang, category, text, file_id])
    wb.save(file_name)

# --- STATES ---
class Form(StatesGroup):
    language = State()
    phone = State()
    category = State()
    content = State()

# --- KEYBOARDS ---
lang_kb = ReplyKeyboardMarkup(
    keyboard=[[KeyboardButton(text="🇺🇿 O'zbek"), KeyboardButton(text="🇷🇺 Русский")]],
    resize_keyboard=True
)

category_uz = ReplyKeyboardMarkup(
    keyboard=[
        [KeyboardButton(text="📝 Text"), KeyboardButton(text="🎤 Voice")],
        [KeyboardButton(text="🎥 Video"), KeyboardButton(text="🖼 Photo")]
    ],
    resize_keyboard=True
)

category_ru = ReplyKeyboardMarkup(
    keyboard=[
        [KeyboardButton(text="📝 Текст"), KeyboardButton(text="🎤 Голос")],
        [KeyboardButton(text="🎥 Видео"), KeyboardButton(text="🖼 Фото")]
    ],
    resize_keyboard=True
)

# --- START ---
@dp.message(CommandStart())
async def start(message: Message, state: FSMContext):
    await state.set_state(Form.language)
    await message.answer("Tilni tanlang / Выберите язык:", reply_markup=lang_kb)

# --- LANGUAGE ---
@dp.message(Form.language)
async def set_lang(message: Message, state: FSMContext):
    lang = "uz" if "O'zbek" in message.text else "ru"
    await state.update_data(lang=lang)
    await state.set_state(Form.phone)

    msg = "📱 Telefon kiriting:" if lang == "uz" else "📱 Введите номер:"
    await message.answer(msg, reply_markup=ReplyKeyboardRemove())

# --- PHONE ---
@dp.message(Form.phone)
async def phone_handler(message: Message, state: FSMContext):
    await state.update_data(phone=message.text)
    data = await state.get_data()

    await state.set_state(Form.category)
    kb = category_uz if data['lang'] == "uz" else category_ru
    msg = "Kategoriya tanlang:" if data['lang'] == "uz" else "Выберите категорию:"

    await message.answer(msg, reply_markup=kb)

# --- CATEGORY ---
@dp.message(Form.category)
async def category_handler(message: Message, state: FSMContext):
    await state.update_data(category=message.text)
    await state.set_state(Form.content)

    await message.answer("Xabar yuboring:", reply_markup=ReplyKeyboardRemove())

# --- CONTENT ---

# --- CONTENT ---
@dp.message(Form.content)
async def content_handler(message: Message, state: FSMContext):
    data = await state.get_data()
    user = message.from_user
    category = data['category']
    
    text = "-"
    file_id = "-"
    valid = False  # Xabar turi to'g'riligini tekshirish uchun

    # Kategoriyaga qarab xabar turini tekshiramiz
    if ("Video" in category or "Видео" in category) and message.video:
        file_id = message.video.file_id
        valid = True
    elif ("Photo" in category or "Фото" in category) and message.photo:
        file_id = message.photo[-1].file_id
        valid = True
    elif ("Voice" in category or "Голос" in category) and message.voice:
        file_id = message.voice.file_id
        valid = True
    elif ("Text" in category or "Текст" in category) and message.text:
        text = message.text
        valid = True

    # Agar tanlangan kategoriya va yuborilgan xabar mos kelmasa
    if not valid:
        msg = f"⚠️ Iltimos, {category} turidagi xabar yuboring!"
        await message.answer(msg)
        return  # Funksiyani to'xtatish, Excelga yozmaydi

    # Agar hammasi to'g'ri bo'lsa, Excelga saqlash
    save_to_excel(user.id, user.username, data['phone'], data['lang'], category, text, file_id)

    admin_msg = f"📥 Yangi xabar\n👤 {user.full_name}\n🆔 {user.id}\n📞 {data['phone']}\n🌍 {data['lang']}\n📌 {category}\n📝 {text if text != '-' else 'Media'}"
    sent = await bot.send_message(ADMIN_ID, admin_msg)
    msg_map[sent.message_id] = user.id

    if file_id != "-":
        if message.photo:
            await bot.send_photo(ADMIN_ID, file_id)
        elif message.video:
            await bot.send_video(ADMIN_ID, file_id)
        elif message.voice:
            await bot.send_voice(ADMIN_ID, file_id)

    await message.answer("✅ Qabul qilindi")
    await state.clear()

# @dp.message(Form.content)
# async def content_handler(message: Message, state: FSMContext):
#     data = await state.get_data()
#     user = message.from_user

#     category = data['category']
#     text = "-"
#     file_id = "-"

#     if message.text:
#         text = message.text
#     elif message.photo:
#         file_id = message.photo[-1].file_id
#     elif message.video:
#         file_id = message.video.file_id
#     elif message.voice:
#         file_id = message.voice.file_id

#     save_to_excel(user.id, user.username, data['phone'], data['lang'], category, text, file_id)






    admin_msg = f"""
📥 Yangi xabar
👤 {user.full_name}
🆔 {user.id}
📞 {data['phone']}
🌍 {data['lang']}
📌 {category}
📝 {text if text != '-' else 'Media'}
"""

    sent = await bot.send_message(ADMIN_ID, admin_msg)
    msg_map[sent.message_id] = user.id

    # media yuborish
    if file_id != "-":
        if message.photo:
            await bot.send_photo(ADMIN_ID, file_id)
        elif message.video:
            await bot.send_video(ADMIN_ID, file_id)
        elif message.voice:
            await bot.send_voice(ADMIN_ID, file_id)

    await message.answer("✅ Qabul qilindi")
    await state.clear()

# --- ADMIN REPLY ---
@dp.message()
async def admin_reply(message: Message):
    if message.from_user.id != ADMIN_ID:
        return

    if not message.reply_to_message:
        return

    user_id = msg_map.get(message.reply_to_message.message_id)
    if not user_id:
        return

    if message.text:
        await bot.send_message(user_id, f"👨‍💻 Admin:\n{message.text}")
    elif message.photo:
        await bot.send_photo(user_id, message.photo[-1].file_id)
    elif message.video:
        await bot.send_video(user_id, message.video.file_id)
    elif message.voice:
        await bot.send_voice(user_id, message.voice.file_id)

import nest_asyncio
import asyncio

# --- RUN ---
async def main():
    print("🚀 Bot hozirgina ishga tushdi...")
    try:
        # Avvalgi sessiyalarni tozalash (xato bermasligi uchun)
        await bot.delete_webhook(drop_pending_updates=True)
        await dp.start_polling(bot)
    except Exception as e:
        print(f"❌ Xatolik yuz berdi: {e}")
    finally:
        await bot.session.close()

if __name__ == "__main__":
    nest_asyncio.apply() # Bu Colab-dagi cheklovni olib tashlaydi
    loop = asyncio.get_event_loop()
    loop.run_until_complete(main())
    

async def main():
    # Mana shu qatorni qo'shing:
    await bot.delete_webhook(drop_pending_updates=True) 
    
    print("✅ Bot muvaffaqiyatli ishga tushdi!")
    await dp.start_polling(bot)
